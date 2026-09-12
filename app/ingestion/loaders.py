"""Document loaders.

Turns files on disk (PDF, DOCX, Markdown, plain text, HTML) into
:class:`Document` objects carrying both the raw text and source metadata that
later ends up in citations.

Loaders are deliberately defensive: a malformed file raises
:class:`DocumentLoadError` with the offending path rather than leaking a
library-specific exception, so the ingestion pipeline can skip and report it.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List

logger = logging.getLogger(__name__)

#: File extensions the loader dispatch table understands.
SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(
    {".pdf", ".txt", ".md", ".markdown", ".docx", ".html", ".htm"}
)


class DocumentLoadError(RuntimeError):
    """Raised when a file cannot be read or parsed."""


@dataclass
class Document:
    """A single logical unit of source text.

    Attributes:
        text: Extracted plain text.
        source: Human-readable source name used in citations (usually the
            file name).
        metadata: Extra provenance such as ``page`` or the absolute path.
    """

    text: str
    source: str
    metadata: Dict[str, object] = field(default_factory=dict)

    @property
    def char_count(self) -> int:
        """Number of characters of extracted text."""
        return len(self.text)

    def is_empty(self) -> bool:
        """True when the document carries no usable text."""
        return not self.text.strip()


def load_from_text(text: str, source_name: str) -> List[Document]:
    """Wrap a raw string as a single :class:`Document`.

    Args:
        text: The text content.
        source_name: Label used for citations, e.g. ``"pasted-note.txt"``.

    Returns:
        A one-element list, matching the shape of the file loaders.
    """
    return [Document(text=text, source=source_name, metadata={"type": "text"})]


def _load_pdf(path: Path) -> List[Document]:
    """Extract one :class:`Document` per PDF page, preserving page numbers."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentLoadError("pypdf is required to read PDF files") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001 - surfaced as a domain error
        raise DocumentLoadError(f"Could not open PDF {path.name}: {exc}") from exc

    documents: List[Document] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001 - one bad page must not kill the file
            logger.warning(
                "Skipping unreadable PDF page", extra={"file": path.name, "page": page_number, "error": str(exc)}
            )
            continue
        if page_text.strip():
            documents.append(
                Document(
                    text=page_text,
                    source=path.name,
                    metadata={
                        "type": "pdf",
                        "page": page_number,
                        "total_pages": len(reader.pages),
                        "path": str(path),
                    },
                )
            )
    return documents


def _load_text(path: Path) -> List[Document]:
    """Read a UTF-8 text or Markdown file, tolerating stray byte sequences."""
    text = path.read_text(encoding="utf-8", errors="replace")
    return [
        Document(
            text=text,
            source=path.name,
            metadata={"type": path.suffix.lstrip(".") or "txt", "path": str(path)},
        )
    ]


def _load_docx(path: Path) -> List[Document]:
    """Extract paragraph text from a Word document."""
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentLoadError("python-docx is required to read .docx files") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 - surfaced as a domain error
        raise DocumentLoadError(f"Could not open DOCX {path.name}: {exc}") from exc

    paragraphs: Iterable[str] = (para.text for para in document.paragraphs)
    text = "\n".join(paragraphs)
    return [
        Document(
            text=text,
            source=path.name,
            metadata={"type": "docx", "path": str(path)},
        )
    ]


def _load_html(path: Path) -> List[Document]:
    """Strip markup from an HTML file, keeping visible text only."""
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentLoadError("beautifulsoup4 is required to read HTML files") from exc

    raw = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    # Collapse the runs of blank lines that markup stripping leaves behind.
    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    return [
        Document(
            text=cleaned,
            source=path.name,
            metadata={"type": "html", "path": str(path)},
        )
    ]


_DISPATCH = {
    ".pdf": _load_pdf,
    ".txt": _load_text,
    ".md": _load_text,
    ".markdown": _load_text,
    ".docx": _load_docx,
    ".html": _load_html,
    ".htm": _load_html,
}


def load_document(path: str | Path) -> List[Document]:
    """Load a single file into one or more documents.

    Args:
        path: Path to the file.

    Returns:
        Extracted documents; empty when the file holds no usable text.

    Raises:
        DocumentLoadError: If the file is missing, unsupported, or unreadable.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise DocumentLoadError(f"File not found: {file_path}")
    if not file_path.is_file():
        raise DocumentLoadError(f"Not a regular file: {file_path}")

    suffix = file_path.suffix.lower()
    loader = _DISPATCH.get(suffix)
    if loader is None:
        raise DocumentLoadError(
            f"Unsupported file type '{suffix}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    documents = [doc for doc in loader(file_path) if not doc.is_empty()]
    logger.info(
        "Loaded document",
        extra={"file": file_path.name, "segments": len(documents), "type": suffix},
    )
    return documents


def load_directory(
    directory: str | Path, recursive: bool = True
) -> tuple[List[Document], List[str]]:
    """Load every supported file in a directory.

    Args:
        directory: Folder to scan.
        recursive: Descend into sub-directories when ``True``.

    Returns:
        A ``(documents, failures)`` tuple.  ``failures`` holds one
        human-readable message per file that could not be read; a single bad
        file never aborts the whole scan.
    """
    root = Path(directory)
    if not root.is_dir():
        raise DocumentLoadError(f"Not a directory: {root}")

    pattern = "**/*" if recursive else "*"
    documents: List[Document] = []
    failures: List[str] = []

    for candidate in sorted(root.glob(pattern)):
        if not candidate.is_file() or candidate.suffix.lower() not in _DISPATCH:
            continue
        try:
            documents.extend(load_document(candidate))
        except DocumentLoadError as exc:
            logger.warning("Skipping document", extra={"error": str(exc)})
            failures.append(str(exc))

    return documents, failures
